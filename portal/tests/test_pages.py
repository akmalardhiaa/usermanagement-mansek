"""The pages HC works in, the status page, and what keeps sessions honest."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.domain.recheck import recheck_sessions
from app.domain.sync import sync_employees
from app.main import create_app
from app.settings import Settings

from conftest import ORIGIN, env_for
from fake_directory import DIRECTORY_ENV, PEOPLE, FakeDirectory


def sign_in(client: TestClient, username: str, password: str) -> TestClient:
    assert client.post("/login", data={"username": username, "password": password}, headers=ORIGIN).status_code == 303
    return client


@pytest.fixture
def demo(make_client):
    """The laptop demo: demo roster, demo sign-in, no directory."""
    return sign_in(make_client(LOGIN_2FA="off"), "hc", "hc12345678")


@pytest.fixture
def office(tmp_path):
    directory = FakeDirectory()
    directory.add("ayu.prameswari", "Ayu Prameswari", "ayu.prameswari@example.com", "Human Capital",
                  f"OU=Human Capital,{PEOPLE}")
    directory.add("hc", "Ayu Prameswari", "", "Human Capital", f"OU=Human Capital,{PEOPLE}")
    directory.add("ops", "Operator TI", "ops@example.com", "IT — Support", f"OU=IT — Support,{PEOPLE}")
    env = env_for(tmp_path, LOGIN_2FA="off", DEMO_LOGIN="off", **DIRECTORY_ENV)
    app = create_app(Settings.from_env(env), directory, env=env, start_background=False)
    return app, directory, env


def test_the_directory_searches_filters_and_pages(demo):
    page = demo.get("/")
    assert page.status_code == 200 and "Rizky Maulana" in page.text
    assert 'href="/pengajuan/baru?jenis=onboarding"' in page.text
    found = demo.get("/", params={"q": "rizky"})
    assert "Rizky Maulana" in found.text and "Dimas Anggara" not in found.text
    none = demo.get("/", params={"q": "tidak-ada-orang-ini"})
    assert "Tidak ada karyawan yang cocok" in none.text
    assert demo.get("/", params={"halaman": "abc"}).status_code == 200


def test_an_employee_page_offers_the_three_requests(demo):
    rizky = next(e for e in demo.app.state.ctx.store.read()["employees"] if e["displayName"] == "Rizky Maulana")
    page = demo.get(f"/karyawan/{rizky['id']}")
    assert page.status_code == 200
    for slug in ("mutasi", "profil", "resign"):
        assert f"jenis={slug}&karyawan={rizky['id']}" in page.text
    assert demo.get("/karyawan/tidak-ada").status_code == 404


def test_every_request_form_draws(demo):
    rizky = next(e for e in demo.app.state.ctx.store.read()["employees"] if e["displayName"] == "Rizky Maulana")
    assert "Pilih jenis pengajuan" in demo.get("/pengajuan/baru").text
    for slug in ("onboarding", "mutasi", "resign", "profil"):
        page = demo.get("/pengajuan/baru", params={"jenis": slug, "karyawan": rizky["id"]})
        assert page.status_code == 200, slug
        assert f'name="jenis" value="{slug}"' in page.text
    assert demo.get("/pengajuan").status_code == 200


def test_the_export_is_text_only_and_is_recorded(demo):
    ctx = demo.app.state.ctx
    ctx.store.mutate(lambda data: data["employees"][0].update(displayName="=HYPERLINK(\"http://evil\")"))
    response = demo.get("/direktori/export.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    names = [row[0].value for row in sheet.iter_rows(min_row=2)]
    cells = [row[0] for row in sheet.iter_rows(min_row=2)]
    assert '=HYPERLINK("http://evil")' in names
    assert all(cell.data_type == "s" for cell in cells)
    assert "diekspor ke Excel" in demo.get("/aktivitas").text


def test_pages_follow_the_role(make_client):
    ops = sign_in(make_client(LOGIN_2FA="off"), "ops", "ops1234567")
    denied = ops.get("/direktori/export.xlsx")
    assert denied.status_code == 403 and "Akses ditolak" in denied.text and "Operator" in denied.text
    assert ops.get("/pengajuan/baru").status_code == 403
    assert ops.get("/status").status_code == 200
    assert ops.get("/status-ad").headers["location"] == "/status"  # the old address still leads there
    hc = sign_in(make_client(LOGIN_2FA="off"), "hc", "hc12345678")
    assert hc.get("/status").status_code == 403
    assert 'href="/status"' not in hc.get("/").text


def test_people_without_a_portal_role_cannot_sign_in(make_client):
    client = make_client(LOGIN_2FA="off")
    response = client.post("/login", data={"username": "bagus", "password": "bagus12345"}, headers=ORIGIN)
    assert response.status_code != 303
    assert client.get("/").headers["location"] == "/login"


def test_the_status_page_reads_the_directory_checks(office):
    app, directory, _ = office
    client = sign_in(TestClient(app, base_url="http://testserver", follow_redirects=False), "ops", "ops1234567")
    page = client.get("/status")
    assert page.status_code == 200
    assert "LDAPS port 636" in page.text and "OU karantina" in page.text
    assert "EMAIL_DRIVER=file" in page.text
    run = client.post("/status/jalankan", data={"apa": "email"}, headers=ORIGIN)
    assert run.status_code == 303 and run.headers["location"].startswith("/status?pesan=")


def test_the_status_page_says_when_the_service_is_down(make_client, ad):
    ad.health = lambda: None
    client = sign_in(make_client(LOGIN_2FA="off"), "ops", "ops1234567")
    assert "Layanan AD tidak menjawab" in client.get("/status").text


def test_sync_adds_updates_and_never_removes(office):
    app, directory, env = office
    store = app.state.ctx.store
    report = sync_employees(store, directory, env)
    assert report["created"] == 2 and report["skipped"] == 1  # the account without a mailbox stays out
    ayu = directory.find_by_name("ayu.prameswari")
    directory.set_attributes(ayu["objectGUID"], {"title": "Head of Human Capital"})
    del directory.accounts[directory.find_by_name("ops")["objectGUID"]]
    report = sync_employees(store, directory, env)
    assert report["updated"] == 1 and report["notFound"] == 1
    emails = {e["email"]: e for e in store.read()["employees"]}
    assert emails["ayu.prameswari@example.com"]["jobTitle"] == "Head of Human Capital"
    assert "ops@example.com" in emails


def test_a_disabled_account_loses_its_session(office):
    app, directory, _ = office
    ctx = app.state.ctx
    client = sign_in(TestClient(app, base_url="http://testserver", follow_redirects=False), "ops", "ops1234567")
    assert client.get("/").status_code == 200

    # AD cannot be read: nobody is thrown out.
    directory.down = True
    assert recheck_sessions(ctx.sessions, directory, ctx.security)["revoked"] == []
    directory.down = False
    assert recheck_sessions(ctx.sessions, directory, ctx.security)["revoked"] == []
    assert client.get("/").status_code == 200

    directory.set_enabled(directory.find_by_name("ops")["objectGUID"], False)
    report = recheck_sessions(ctx.sessions, directory, ctx.security)
    assert [r["username"] for r in report["revoked"]] == ["ops"]
    assert client.get("/").headers["location"] == "/login"
