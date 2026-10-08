"""The employee directory: search, filter, sort, one page at a time, and export.

Filtering and sorting run here, on the server, over everybody; the page draws
fifty rows. A directory read from AD with a thousand people stays quick.
"""

from __future__ import annotations

import io

from fastapi import FastAPI, Request
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.domain.catalog import DEPARTMENT_GROUPS, DEPARTMENTS, EMPLOYMENT_TYPES, LOCATIONS
from app.domain.lifecycle import is_active
from app.domain.store import now_iso, record_activity
from app.web.context import Context

PAGE_SIZE = 50
# Sort keys as the address carries them. "diperbarui" is the default: sorted by
# name, a hire the worker had just created landed mid-roster and HC, looking at
# the top for the person they had just onboarded, concluded it had not worked.
SORTS = {"diperbarui": "updatedAt", "nama": "displayName", "jabatan": "jobTitle", "departemen": "department",
         "status": "status"}
SORT_LABELS = {"diperbarui": "Terakhir diperbarui", "nama": "Nama", "jabatan": "Jabatan", "departemen": "Departemen",
               "status": "Status"}
STATUS_FILTERS = {"": "Semua status", "pengajuan": "Dalam persetujuan", "aktif": "Aktif", "nonaktif": "Nonaktif"}


def _query(params) -> dict:
    """The filters from the address, each coerced to a value the page can draw."""
    status = (params.get("status") or "").strip()
    status = "" if status == "semua" else status
    sort = params.get("urut") or "diperbarui"
    return {
        "q": (params.get("q") or "").strip()[:200],
        "status": status if status in STATUS_FILTERS else "",
        "departemen": (params.get("departemen") or "").strip()[:200],
        "urut": sort if sort in SORTS else "diperbarui",
        "arah": "turun" if params.get("arah") == "turun" else "naik",
    }


def _filtered(employees: list[dict], query: dict, pending_ids: set[str] | frozenset = frozenset()) -> list[dict]:
    text = query["q"].lower()
    status = query["status"]
    found = []
    for employee in employees:
        if status == "aktif" and employee["status"] != "ACTIVE":
            continue
        if status == "nonaktif" and employee["status"] != "DISABLED":
            continue
        if status == "pengajuan" and employee["id"] not in pending_ids:
            continue
        # The first view shows working accounts; a disabled one appears when
        # searched for by name or asked for by status.
        if status == "" and not text and employee["status"] != "ACTIVE":
            continue
        if query["departemen"] and employee.get("department") != query["departemen"]:
            continue
        if text and not any(text in (employee.get(f) or "").lower() for f in ("displayName", "email", "jobTitle", "department")):
            continue
        found.append(employee)
    key = SORTS[query["urut"]]
    if key == "updatedAt":
        # Newest first is the natural "up" for a date.
        ordered = sorted(found, key=lambda e: e.get("updatedAt") or e.get("createdAt") or "", reverse=True)
        return list(reversed(ordered)) if query["arah"] == "turun" else ordered
    return sorted(found, key=lambda e: ((e.get(key) or "").lower(), e["displayName"].lower()),
                  reverse=query["arah"] == "turun")


def _department_groups(present: set[str]) -> list[tuple[str, list[str]]]:
    """The catalogue's headings, then whatever the directory holds that the catalogue does not."""
    groups = [(label, [d for d in items if d in present]) for label, items in DEPARTMENT_GROUPS]
    groups = [(label, items) for label, items in groups if items]
    known = {d for _, items in groups for d in items}
    other = sorted(present - known, key=str.lower)
    return groups + ([("Lainnya", other)] if other else [])


PIPELINE_STEPS = ["Diajukan", "Manager", "CISO", "Dibuat di AD"]


def _pipeline_item(request: dict) -> dict:
    """Where an onboarding stands, for people who have no account yet."""
    status, payload = request["status"], request["payload"]
    manager = next((s["approver"]["name"] for s in request.get("approvals", [])
                    if s["stage"] == "MANAGER" and s["version"] == request["version"]), payload.get("managerName", ""))
    step, tone, label = {
        "DRAFT": (0, "muted", "Draf \u2014 belum dikirim ke approver"),
        "PENDING_MANAGER": (1, "warn", f"Menunggu persetujuan {manager}"),
        "PENDING_CISO": (2, "warn", "Menunggu persetujuan tim CISO"),
        "EXECUTING": (3, "info", "Sedang dibuat di AD\u2026"),
        "FAILED": (3, "danger", "Gagal dibuat di AD \u2014 buka detail untuk dicek"),
    }.get(status, (3, "ok", "Disetujui \u2014 siap dibuat di AD"))
    return {"id": request["id"], "name": payload.get("displayName", request["subject"]["displayName"]),
            "jobTitle": payload.get("jobTitle", ""), "department": payload.get("department", ""),
            "email": payload.get("email", ""), "step": step, "tone": tone, "label": label,
            "ready": status in ("QUEUED", "APPROVED", "SCHEDULED")}


def register(app: FastAPI, ctx: Context) -> None:
    @app.get("/", response_class=HTMLResponse)
    def directory(request: Request):
        session, refused = ctx.guard(request, "directory.read")
        if refused:
            return refused
        data = ctx.store.read()
        query = _query(request.query_params)
        active_requests = {r["employeeId"]: r for r in data["lifecycleRequests"]
                           if r.get("employeeId") and is_active(r["status"])}
        employees = _filtered(data["employees"], query, set(active_requests))
        try:
            page = max(1, int(request.query_params.get("halaman") or 1))
        except ValueError:
            page = 1
        pages = max(1, -(-len(employees) // PAGE_SIZE))
        page = min(page, pages)
        stats = {
            "total": len(data["employees"]),
            "active": sum(1 for e in data["employees"] if e["status"] == "ACTIVE"),
            "pending": len(active_requests),
            "disabled": sum(1 for e in data["employees"] if e["status"] == "DISABLED"),
        }
        hidden_inactive = (stats["total"] - stats["active"]) if not query["status"] and not query["q"] else 0
        onboardings = [_pipeline_item(r) for r in sorted(data["lifecycleRequests"], key=lambda r: r["createdAt"], reverse=True)
                       if r["type"] == "ONBOARDING" and is_active(r["status"])]
        base_query = {k: v for k, v in query.items()
                      if v and not (k == "urut" and v == "diperbarui") and not (k == "arah" and v == "naik")}
        present = {e["department"] for e in data["employees"] if e.get("department")}
        return ctx.page(request, "directory.html", session=session, wide=True,
                        employees=employees[(page - 1) * PAGE_SIZE: page * PAGE_SIZE], total=len(employees),
                        everybody=stats["total"], page=page, pages=pages, query=query, base_query=base_query,
                        department_groups=_department_groups(present | set(DEPARTMENTS)), statuses=STATUS_FILTERS,
                        sort_labels=SORT_LABELS, stats=stats, active_requests=active_requests,
                        hidden_inactive=hidden_inactive, onboardings=onboardings, steps=PIPELINE_STEPS,
                        dirty=bool(query["q"] or query["status"] or query["departemen"]),
                        message=request.query_params.get("pesan"))

    @app.get("/karyawan/{employee_id}", response_class=HTMLResponse)
    def employee_page(request: Request, employee_id: str):
        session, refused = ctx.guard(request, "directory.read")
        if refused:
            return refused
        data = ctx.store.read()
        employee = next((e for e in data["employees"] if e["id"] == employee_id), None)
        if employee is None:
            return ctx.page(request, "not_found.html", session=session, status=404)
        history = sorted((r for r in data["lifecycleRequests"] if r.get("employeeId") == employee_id),
                         key=lambda r: r["createdAt"], reverse=True)
        active = next((r for r in history if is_active(r["status"])), None)
        return ctx.page(request, "employee.html", session=session, employee=employee, history=history, active=active,
                        employment=EMPLOYMENT_TYPES, locations=LOCATIONS)

    @app.get("/direktori/export.xlsx")
    def export(request: Request):
        session, refused = ctx.guard(request, "directory.export")
        if refused:
            return refused
        data = ctx.store.read()
        pending_ids = {r["employeeId"] for r in data["lifecycleRequests"] if r.get("employeeId") and is_active(r["status"])}
        employees = _filtered(data["employees"], _query(request.query_params), pending_ids)
        book = Workbook()
        sheet = book.active
        sheet.title = "Direktori"
        headers = ["Nama", "Email", "Jabatan", "Departemen", "Status", "Status kepegawaian", "Lokasi", "Manager", "Email manager"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="173D6E")
        for e in employees:
            sheet.append([e["displayName"], e["email"], e.get("jobTitle", ""), e.get("department", ""),
                          "Aktif" if e["status"] == "ACTIVE" else "Nonaktif",
                          EMPLOYMENT_TYPES.get(e.get("employmentType") or "", ""), LOCATIONS.get(e.get("locationType") or "", ""),
                          e.get("managerName", ""), e.get("managerEmail", "")])
            # Text stays text: a name from AD that starts with "=" must not become a formula.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for column, width in zip("ABCDEFGHI", (28, 34, 26, 22, 10, 20, 14, 24, 32)):
            sheet.column_dimensions[column].width = width
        buffer = io.BytesIO()
        book.save(buffer)
        buffer.seek(0)
        ctx.store.mutate(lambda data: record_activity(
            data, session.full_name, "directory.exported", f"{len(employees)} baris diekspor ke Excel."))
        name = f"direktori-karyawan-{now_iso()[:10]}.xlsx"
        return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="{name}"'})
